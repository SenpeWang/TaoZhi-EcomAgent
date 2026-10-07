import io,zipfile,hashlib
from .policy import Denied
def validate_upload(name,data):
    if not data or len(data)>20*1024*1024:raise Denied("附件大小须为 1 字节至 20 MB","UPLOAD_SIZE",400)
    suffix=name.lower().rsplit(".",1)[-1]
    if suffix=="pdf":
        if not data.startswith(b"%PDF-"):raise Denied("PDF 格式无效","UPLOAD_FORMAT",400)
        return "application/pdf"
    if suffix=="docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist() or sum(x.file_size for x in z.infolist())>30*1024*1024 or any(x.filename.startswith("/") or ".." in x.filename.split("/") for x in z.infolist()):raise ValueError()
        except Exception:raise Denied("文档格式无效或解压内容过大","UPLOAD_FORMAT",400) from None
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if suffix in ("txt","md"):
        try:data.decode("utf-8-sig")
        except UnicodeError:raise Denied("文本附件须为 UTF-8 编码","UPLOAD_ENCODING",400) from None
        return "text/plain"
    raise Denied("支持 PDF、Word、文本和 Markdown 附件","UPLOAD_FORMAT",400)
def parse(path,mime):
    data=path.read_bytes()
    if mime=="application/pdf":
        import fitz
        with fitz.open(stream=data,filetype="pdf") as pdf:
            if pdf.is_encrypted or len(pdf)>200:raise Denied("暂不支持加密或超过 200 页的 PDF","PARSE_LIMIT",400)
            pages=[(i+1,pdf[i].get_text()) for i in range(len(pdf))]
    elif "wordprocessingml" in mime:
        from docx import Document
        doc=Document(io.BytesIO(data))
        pages=[(0,"\n".join([p.text for p in doc.paragraphs]+[" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]))]
    else:pages=[(0,data.decode("utf-8-sig"))]
    body="\n\n".join(text for _,text in pages)
    if not body.strip() or len(body)>200000:raise Denied("未提取到可读文本或内容超过限制","PARSE_LIMIT",400)
    return body,pages
def split(pages,doc_id,version):
    result=[]
    for page,text in pages:
        for start in range(0,len(text),650):
            value=text[start:start+850].strip()
            if value:result.append(dict(id="chunk_"+hashlib.sha256(f"{doc_id}:{version}:{len(result)}".encode()).hexdigest()[:24],ordinal=len(result),page=page,text=value))
    return result
