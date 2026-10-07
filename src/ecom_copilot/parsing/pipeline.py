"""异步解析 Pipeline（PRD B5）。

链路：MinIO 分片上传 → Redis BitMap 记录分片上报状态 → Kafka 异步承接解析与向量化任务。
任一组件不可用时自动降级为本地实现，接口保持一致。
"""

from __future__ import annotations

import json
import threading
import time
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional

from ..config import Settings, get_settings
from ..schemas.common import DocType, SourceDocument
from ..storage import content_hash, get_blob_store, get_chunk_state, get_event_bus
from .chunker import chunk_document
from .pdf import parse_pdf_file, tables_to_text
from .vision import VisionParser

PART_SIZE = 4 * 1024 * 1024


class ParseJob:
    def __init__(self, upload_id: str, filename: str, total_parts: int) -> None:
        self.upload_id = upload_id
        self.filename = filename
        self.total_parts = total_parts
        self.status = "uploading"     # uploading | queued | parsing | done | failed
        self.progress = 0.0
        self.document: Optional[SourceDocument] = None
        self.chunks: List[Any] = []
        self.error = ""
        self.created_at = time.time()
        self.updated_at = time.time()

    def as_dict(self) -> Dict[str, Any]:
        return {
            "upload_id": self.upload_id,
            "filename": self.filename,
            "total_parts": self.total_parts,
            "status": self.status,
            "progress": round(self.progress, 3),
            "error": self.error,
            "doc_id": self.document.id if self.document else "",
            "chunks": len(self.chunks),
        }


class DocumentParsingPipeline:
    """分片上传 + 断点续传 + 异步解析。"""

    def __init__(self, settings: Optional[Settings] = None) -> None:
        self.settings = settings or get_settings()
        self.blobs = get_blob_store(self.settings)
        self.state = get_chunk_state(self.settings)
        self.bus = get_event_bus(self.settings)
        self.vision = VisionParser(self.settings)
        self._jobs: Dict[str, ParseJob] = {}
        self._lock = threading.RLock()
        self.work_dir = Path(self.settings.data_dir) / "parsing"
        self.work_dir.mkdir(parents=True, exist_ok=True)
        try:
            self.bus.subscribe(self.settings.kafka_topic_parse, self._handle_parse_event)
        except Exception:  # noqa: BLE001
            pass

    # ─────────────── 分片上传 ───────────────
    def init_upload(self, filename: str, size: int) -> Dict[str, Any]:
        upload_id = f"{content_hash(filename.encode('utf-8'))}-{uuid.uuid4().hex[:8]}"
        total_parts = max(1, (size + PART_SIZE - 1) // PART_SIZE)
        with self._lock:
            self._jobs[upload_id] = ParseJob(upload_id, filename, total_parts)
        return {"upload_id": upload_id, "total_parts": total_parts, "part_size": PART_SIZE}

    def upload_part(self, upload_id: str, index: int, data: bytes) -> Dict[str, Any]:
        job = self._jobs.get(upload_id)
        if not job:
            raise KeyError(f"未知 upload_id: {upload_id}")
        key = f"{upload_id}/part-{index:05d}"
        self.blobs.put(key, data)
        self.state.mark(upload_id, index, job.total_parts)
        done = self.state.completed(upload_id)
        job.progress = len(done) / job.total_parts
        if self.state.is_complete(upload_id, job.total_parts):
            job.status = "queued"
            self.bus.publish(
                self.settings.kafka_topic_parse,
                {"upload_id": upload_id, "filename": job.filename, "total_parts": job.total_parts},
            )
            # 本地总线同步兜底：确保无 Kafka 时也能推进
            self._handle_parse_event({"upload_id": upload_id})
        return {"upload_id": upload_id, "done": len(done), "total": job.total_parts,
                "progress": round(job.progress, 3), "status": job.status}

    def upload_progress(self, upload_id: str) -> Dict[str, Any]:
        job = self._jobs.get(upload_id)
        if not job:
            return {}
        return job.as_dict()

    # ─────────────── 一步提交 ───────────────
    def submit_bytes(self, data: bytes, filename: str,
                     doc_type: DocType = DocType.MANUAL) -> SourceDocument:
        info = self.init_upload(filename,len(data))
        upload_id = info["upload_id"]
        # Preserve part bookkeeping without publishing the same parse twice.
        for index in range(info["total_parts"]):
            self.blobs.put(f"{upload_id}/part-{index:05d}",data[index*PART_SIZE:(index+1)*PART_SIZE])
            self.state.mark(upload_id,index,info["total_parts"])
        doc = self._parse_and_store(upload_id, filename, data, doc_type)
        job = self._jobs[upload_id]
        job.status = "done"
        job.document = doc
        job.chunks = chunk_document(doc)
        job.progress = 1.0
        return doc

    # ─────────────── 解析 ───────────────
    def _handle_parse_event(self, event: Dict[str, Any]) -> None:
        upload_id = event.get("upload_id")
        if not upload_id:
            return
        job = self._jobs.get(upload_id)
        if not job or job.status in ("parsing", "done"):
            return
        job.status = "parsing"
        data = self._reassemble(upload_id, job.total_parts)
        if data is None:
            job.status = "failed"
            job.error = "分片缺失，无法重组文件"
            return
        try:
            doc = self._parse_and_store(upload_id, job.filename, data, DocType.MANUAL)
            job.document = doc
            job.chunks = chunk_document(doc)
            job.status = "done"
            job.progress = 1.0
        except Exception as exc:  # noqa: BLE001
            job.status = "failed"
            job.error = f"{type(exc).__name__}: {exc}"
        job.updated_at = time.time()

    def _reassemble(self, upload_id: str, total_parts: int) -> Optional[bytes]:
        blobs: List[bytes] = []
        for index in range(total_parts):
            part = self.blobs.get(f"{upload_id}/part-{index:05d}")
            if part is None:
                return None
            blobs.append(part)
        return b"".join(blobs)

    def _parse_and_store(self, upload_id: str, filename: str, data: bytes,
                         doc_type: DocType) -> SourceDocument:
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            path = self.work_dir / f"{upload_id}.pdf"
            path.write_bytes(data)
            doc = parse_pdf_file(path, doc_type=doc_type, title=Path(filename).stem)
        else:
            if suffix == ".docx":
                import io
                from docx import Document
                document=Document(io.BytesIO(data))
                text="\n".join([p.text for p in document.paragraphs]+[
                    " | ".join(c.text for c in row.cells)
                    for table in document.tables for row in table.rows])
            else:
                text = data.decode("utf-8", errors="strict")
            doc = SourceDocument(
                title=Path(filename).stem,
                source="upload",
                doc_type=doc_type,
                raw_text=text,
                url=f"file://{filename}",
            )
        if doc.tables:
            doc.tables_description = tables_to_text(doc.tables)
        if doc.images and self.vision.enabled:
            descriptions = self.vision.describe_batch(doc.images, limit=3)
            if descriptions:
                doc.tables_description = (
                    doc.tables_description + "\n\n【图表解析】\n" + "\n".join(descriptions)
                ).strip()
        (self.work_dir / f"{upload_id}.json").write_text(
            json.dumps(doc.model_dump(mode="json"), ensure_ascii=False, default=str),
            encoding="utf-8",
        )
        return doc

    def get_document(self, upload_id: str) -> Optional[SourceDocument]:
        job = self._jobs.get(upload_id)
        if job and job.document:
            return job.document
        path = self.work_dir / f"{upload_id}.json"
        if path.exists():
            return SourceDocument.model_validate_json(path.read_text(encoding="utf-8"))
        return None

    def wait_for(self, upload_id: str, timeout: float = 60.0) -> Optional[SourceDocument]:
        deadline = time.time() + timeout
        while time.time() < deadline:
            doc = self.get_document(upload_id)
            if doc is not None:
                return doc
            time.sleep(0.3)
        return None


_pipeline: Optional[DocumentParsingPipeline] = None


def get_pipeline(settings: Optional[Settings] = None) -> DocumentParsingPipeline:
    global _pipeline
    if _pipeline is None:
        _pipeline = DocumentParsingPipeline(settings)
    return _pipeline
