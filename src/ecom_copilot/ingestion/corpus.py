"""企业语料文本抽取。

只做一件事：把一个 md/txt/pdf/docx 文件读成纯文本。原来放在 sources.py 里，
而 sources.py 依赖已下线的抓取适配器与 schemas；评估链路（evals/dataset）
要用这个抽取能力，所以单独成模块，不牵连任何已删除的东西。
"""
from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

_TEXT_SUFFIXES = {".md", ".txt", ".markdown"}
_PDF_SUFFIXES = {".pdf"}
_DOCX_SUFFIXES = {".docx"}


def read_corpus_text(path: Path) -> str:
    """读取文档文本：md/txt 直接读取；pdf/docx 按可用库抽取，全部失败才留空。"""
    suffix = path.suffix.lower()
    if suffix in _PDF_SUFFIXES:
        # PyMuPDF 优先：pypdf 在本环境未安装，只写 pypdf 会让 PDF 静默解析成空串，
        # 表现为"文档入库了但检索永远命中不到"，排查成本极高。
        try:
            import pymupdf  # noqa: PLC0415

            with pymupdf.open(str(path)) as doc:
                return "\n".join(page.get_text("text") for page in doc)
        except Exception as exc:  # noqa: BLE001
            logger.warning("PyMuPDF 抽取失败，尝试 pypdf: %s (%s)", path, exc)
        try:
            from pypdf import PdfReader  # noqa: PLC0415

            return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
        except Exception as exc:  # noqa: BLE001
            logger.warning("PDF 文本抽取失败，已跳过内容: %s (%s)", path, exc)
            return ""
    if suffix in _DOCX_SUFFIXES:
        try:
            from docx import Document  # noqa: PLC0415

            d = Document(str(path))
            parts = [p.text for p in d.paragraphs]
            # 页眉页脚不属于 body：漏读会让版本号、备案号、生效日期整段丢失
            for section in d.sections:
                parts += [p.text for p in section.header.paragraphs]
                parts += [p.text for p in section.footer.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(x for x in parts if x.strip())
        except Exception as exc:  # noqa: BLE001
            logger.warning("DOCX 文本抽取失败，已跳过内容: %s (%s)", path, exc)
            return ""
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except OSError as exc:
        logger.warning("文档读取失败，已跳过: %s (%s)", path, exc)
        return ""
