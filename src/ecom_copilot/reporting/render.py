"""报告多格式导出：Markdown / Word / PDF / HTML（PRD F4）。"""

from __future__ import annotations

import html
from pathlib import Path
from typing import Any, Dict

from ..config import Settings, get_settings
from ..schemas.research import ResearchReport


def to_markdown(report: ResearchReport) -> str:
    return report.to_markdown()


def to_html(report: ResearchReport) -> str:
    """自包含 HTML（内联 CSS），可直接浏览器打开。"""
    sections = "".join(
        f"<section><h2>{html.escape(s.heading)}</h2>"
        f"<p>{html.escape(s.content).replace(chr(10), '<br/>')}</p>"
        f"<p class='refs'>引用：{' '.join(str(i) for i in sorted(set(s.citation_indices)))}</p></section>"
        for s in report.sections
    )
    citations = "".join(
        f"<li>[{c.index}] {html.escape(c.doc_name)} "
        f"{'P' + str(c.page) if c.page else '全文'}"
        f"{f' <a href={html.escape(c.url)!r}>原文</a>' if c.url else ''}</li>"
        for c in report.citations
    )
    return f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"/>
<title>{html.escape(report.title)}</title>
<style>
body{{font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif;max-width:900px;margin:32px auto;line-height:1.75;color:#1f2328}}
h1{{border-bottom:2px solid #2f6feb;padding-bottom:8px}}
h2{{color:#2f6feb;margin-top:28px}}
.refs{{color:#6e7781;font-size:13px}}
.meta{{color:#6e7781;font-size:13px}}
section{{margin-bottom:20px}}
</style></head><body>
<h1>{html.escape(report.title)}</h1>
<p class="meta">生成时间：{report.generated_at:%Y-%m-%d %H:%M} ｜ 置信度：{report.confidence:.2f} ｜ 引用覆盖率：{report.citation_coverage:.0%}</p>
<h2>摘要</h2><p>{html.escape(report.executive_summary).replace(chr(10), '<br/>')}</p>
{sections}
<h2>引用来源</h2><ol>{citations}</ol>
</body></html>"""


def to_docx(report: ResearchReport, path: Path) -> Path:
    from docx import Document  # noqa: PLC0415
    from docx.shared import Pt  # noqa: PLC0415

    doc = Document()
    doc.add_heading(report.title, level=0)
    doc.add_paragraph(
        f"生成时间：{report.generated_at:%Y-%m-%d %H:%M}｜置信度：{report.confidence:.2f}｜"
        f"引用溯源覆盖率：{report.citation_coverage:.0%}"
    )
    doc.add_heading("摘要", level=1)
    doc.add_paragraph(report.executive_summary)
    for section in report.sections:
        doc.add_heading(section.heading, level=1)
        doc.add_paragraph(section.content)
        if section.citation_indices:
            refs = doc.add_paragraph(
                "引用：" + " ".join(f"[{i}]" for i in sorted(set(section.citation_indices)))
            )
            refs.runs[0].font.size = Pt(9)
    if report.citations:
        doc.add_heading("引用来源", level=1)
        for citation in report.citations:
            page = f"P{citation.page}" if citation.page else "全文"
            doc.add_paragraph(f"[{citation.index}] {citation.doc_name} {page}")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def to_pdf(report: ResearchReport, path: Path) -> Path:
    """reportlab 中文输出（内置 CID 字体 STSong-Light）。"""
    from reportlab.lib.pagesizes import A4  # noqa: PLC0415
    from reportlab.lib.styles import ParagraphStyle  # noqa: PLC0415
    from reportlab.lib.units import mm  # noqa: PLC0415
    from reportlab.pdfbase import pdfmetrics  # noqa: PLC0415
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # noqa: PLC0415
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # noqa: PLC0415

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = {
        "title": ParagraphStyle("title", fontName="STSong-Light", fontSize=18, leading=24),
        "h": ParagraphStyle("h", fontName="STSong-Light", fontSize=13, leading=18,
                            spaceBefore=10),
        "body": ParagraphStyle("body", fontName="STSong-Light", fontSize=10, leading=15),
    }
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    story = [Paragraph(_esc(report.title), styles["title"]), Spacer(1, 6)]
    story.append(Paragraph(_esc(f"置信度 {report.confidence:.2f}｜引用覆盖率 {report.citation_coverage:.0%}"),
                           styles["body"]))
    story.append(Paragraph("摘要", styles["h"]))
    story.append(Paragraph(_esc(report.executive_summary), styles["body"]))
    for section in report.sections:
        story.append(Paragraph(_esc(section.heading), styles["h"]))
        story.append(Paragraph(_esc(section.content), styles["body"]))
    if report.citations:
        story.append(Paragraph("引用来源", styles["h"]))
        for citation in report.citations:
            page = f"P{citation.page}" if citation.page else "全文"
            story.append(Paragraph(_esc(f"[{citation.index}] {citation.doc_name} {page}"),
                                   styles["body"]))
    doc.build(story)
    return path


def _esc(text: str) -> str:
    return html.escape(str(text or "")).replace("\n", "<br/>")


def export(report: ResearchReport, fmt: str, path: Path,
           settings: Settings | None = None) -> Path:
    settings = settings or get_settings()
    path = Path(path)
    if fmt == "markdown":
        path.write_text(to_markdown(report), encoding="utf-8")
    elif fmt == "html":
        path.write_text(to_html(report), encoding="utf-8")
    elif fmt == "docx":
        to_docx(report, path)
    elif fmt == "pdf":
        to_pdf(report, path)
    else:
        raise ValueError(f"不支持的导出格式：{fmt}")
    return path


def summary_card(report: ResearchReport) -> Dict[str, Any]:
    return {
        "title": report.title,
        "confidence": report.confidence,
        "citation_coverage": report.citation_coverage,
        "field_fill_rate": round(report.field_fill_rate(), 3),
        "sections": len(report.sections),
        "citations": len(report.citations),
        "key_companies": len(report.key_companies),
        "review_required": report.review_required,
    }
